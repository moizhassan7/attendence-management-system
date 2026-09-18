"""Attendance API endpoints — raw punches and daily records."""

from __future__ import annotations

from datetime import date, date as dt_date, datetime, timedelta
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel
from sqlalchemy import select, func, or_
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.attendance import AttendancePunch, AttendanceDaily
from app.models.exception import AttendanceException
from app.models.personnel import Personnel
from app.models.device import Device
from app.models.settings import SystemSetting
from app.models.shift import Shift
from app.schemas.common import ApiResponse, PaginatedResponse, PaginationMeta
from app.schemas.attendance import AttendancePunchOut, AttendanceDailyOut
from app.services.attendance_engine import AttendanceService, infer_shift_for_punch, is_security_staff
from app.services.pin_match import candidate_device_pins, resolve_personnel_pin
from app.utils.timezone import today, now, to_local, get_tz

router = APIRouter(prefix="/attendance", tags=["attendance"])


def _person_for_punch_pin(punch_pin: str, people_by_pin: dict[str, Personnel]):
    resolved = resolve_personnel_pin(punch_pin, people_by_pin)
    if resolved:
        return people_by_pin.get(resolved)
    return people_by_pin.get(str(punch_pin))


def _matches_quick_filter(status: str, quick_filter: str) -> bool:
    if quick_filter in ("all", "by_department", "summary"):
        return True
    if quick_filter == "present":
        return status in ("Present", "Late")
    if quick_filter == "late":
        return status == "Late"
    if quick_filter == "absent":
        return status in ("Absent", "Weekend")
    if quick_filter == "leave":
        return status == "Leave"
    return True


_STATUS_SORT = {"Present": 0, "Late": 1, "Leave": 2, "Weekend": 3, "Absent": 4}


@router.get("/punches", response_model=PaginatedResponse)
async def list_punches(
    page: Annotated[int, Query(ge=1)] = 1,
    page_size: Annotated[int, Query(ge=1, le=1000)] = 50,
    date_filter: Annotated[date | None, Query(alias="date")] = None,
    biometric_user_id: Annotated[str | None, Query()] = None,
    device_id: Annotated[int | None, Query()] = None,
    db: AsyncSession = Depends(get_db),
):
    """List raw attendance punches with filters."""
    query = select(AttendancePunch)
    count_query = select(func.count(AttendancePunch.id))

    if date_filter:
        from datetime import datetime, time
        tz = get_tz()
        t0 = datetime.combine(date_filter, time.min, tzinfo=tz)
        t1 = datetime.combine(date_filter, time.max, tzinfo=tz)
        query = query.where(AttendancePunch.punch_time >= t0, AttendancePunch.punch_time <= t1)
        count_query = count_query.where(AttendancePunch.punch_time >= t0, AttendancePunch.punch_time <= t1)

    if biometric_user_id:
        pins = candidate_device_pins(biometric_user_id)
        query = query.where(AttendancePunch.biometric_user_id.in_(pins))
        count_query = count_query.where(AttendancePunch.biometric_user_id.in_(pins))

    if device_id:
        query = query.where(AttendancePunch.device_id == device_id)
        count_query = count_query.where(AttendancePunch.device_id == device_id)

    total = (await db.execute(count_query)).scalar() or 0
    result = await db.execute(
        query.order_by(AttendancePunch.punch_time.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    punches = result.scalars().all()
    people = (await db.execute(select(Personnel))).scalars().all()
    people_by_pin = {str(person.biometric_user_id): person for person in people}

    out = []
    for p in punches:
        person = _person_for_punch_pin(p.biometric_user_id, people_by_pin)
        out.append(AttendancePunchOut(
            id=p.id,
            device_id=p.device_id,
            biometric_user_id=p.biometric_user_id,
            punch_time=to_local(p.punch_time),
            punch_type=p.punch_type,
            verified=p.verified,
            source=p.source,
            created_at=to_local(p.created_at),
            personnel_name=person.full_name if person else None,
            device_name=p.device.name if p.device else None,
        ))

    return PaginatedResponse(
        data=out,
        pagination=PaginationMeta(page=page, page_size=page_size, total=total),
    )


@router.get("/daily", response_model=PaginatedResponse)
async def list_daily_attendance(
    page: Annotated[int, Query(ge=1)] = 1,
    page_size: Annotated[int, Query(ge=1, le=1000)] = 50,
    date_filter: Annotated[date | None, Query(alias="date")] = None,
    department_id: Annotated[int | None, Query()] = None,
    rank_id: Annotated[int | None, Query()] = None,
    course_id: Annotated[int | None, Query()] = None,
    category: Annotated[str | None, Query()] = None,
    duty_type: Annotated[str | None, Query()] = None,
    is_trainee: Annotated[bool | None, Query()] = None,
    status: Annotated[str | None, Query()] = None,
    search: Annotated[str | None, Query()] = None,
    db: AsyncSession = Depends(get_db),
):
    """List daily attendance records with extensive filters and search."""
    target_date = date_filter or today()

    # If no records exist yet for this date, trigger calculation once
    existing_count = await db.scalar(
        select(func.count(AttendanceDaily.id)).where(AttendanceDaily.attendance_date == target_date)
    )
    if not existing_count:
        try:
            await AttendanceService().process_daily_attendance(db, target_date)
        except Exception as exc:
            logger.warning("Could not pre-process daily attendance for %s: %s", target_date, exc)

    query = select(AttendanceDaily).options(
        selectinload(AttendanceDaily.personnel).selectinload(Personnel.rank),
        selectinload(AttendanceDaily.personnel).selectinload(Personnel.department),
        selectinload(AttendanceDaily.personnel).selectinload(Personnel.course),
        selectinload(AttendanceDaily.personnel).selectinload(Personnel.shift),
    )
    count_query = select(func.count(AttendanceDaily.id))

    query = query.where(AttendanceDaily.attendance_date == target_date)
    count_query = count_query.where(AttendanceDaily.attendance_date == target_date)

    # Status filter (supports comma-separated e.g. "PRESENT,LATE" or single status)
    if status and status.strip().upper() != "ALL":
        if "," in status:
            statuses = [s.strip().upper() for s in status.split(",") if s.strip()]
            query = query.where(func.upper(AttendanceDaily.status).in_(statuses))
            count_query = count_query.where(func.upper(AttendanceDaily.status).in_(statuses))
        else:
            query = query.where(func.upper(AttendanceDaily.status) == status.strip().upper())
            count_query = count_query.where(func.upper(AttendanceDaily.status) == status.strip().upper())

    # Personnel filters
    needs_personnel_join = any([
        department_id is not None,
        rank_id is not None,
        course_id is not None,
        category is not None,
        duty_type is not None,
        is_trainee is not None,
        bool(search and search.strip()),
    ])

    if needs_personnel_join:
        query = query.join(Personnel, AttendanceDaily.personnel_id == Personnel.id)
        count_query = count_query.join(Personnel, AttendanceDaily.personnel_id == Personnel.id)

        if department_id:
            query = query.where(Personnel.department_id == department_id)
            count_query = count_query.where(Personnel.department_id == department_id)
        if rank_id:
            query = query.where(Personnel.rank_id == rank_id)
            count_query = count_query.where(Personnel.rank_id == rank_id)
        if course_id:
            query = query.where(Personnel.course_id == course_id)
            count_query = count_query.where(Personnel.course_id == course_id)
        if category:
            query = query.where(Personnel.category == category)
            count_query = count_query.where(Personnel.category == category)
        if duty_type:
            query = query.where(Personnel.duty_type == duty_type)
            count_query = count_query.where(Personnel.duty_type == duty_type)
        if is_trainee is not None:
            query = query.where(Personnel.is_trainee == is_trainee)
            count_query = count_query.where(Personnel.is_trainee == is_trainee)
        if search and search.strip():
            term = f"%{search.strip()}%"
            s_filter = or_(
                Personnel.full_name.ilike(term),
                Personnel.biometric_user_id.ilike(term),
                Personnel.employee_code.ilike(term),
                Personnel.cnic.ilike(term),
            )
            query = query.where(s_filter)
            count_query = count_query.where(s_filter)

    total = (await db.execute(count_query)).scalar() or 0
    result = await db.execute(
        query.order_by(AttendanceDaily.personnel_id)
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    records = result.scalars().all()

    active_shifts = None
    out = []
    for r in records:
        p = r.personnel
        s_name = p.shift.name if p and p.shift else None
        if not s_name and p and is_security_staff(p.duty_type):
            if r.first_in:
                if active_shifts is None:
                    active_shifts = list((await db.execute(select(Shift).where(Shift.active == True))).scalars().all())
                inferred = infer_shift_for_punch(to_local(r.first_in), active_shifts)
                s_name = inferred.name if inferred else "Morning Shift"
            elif r.status in ["LEAVE", "MEDICAL", "OSD", "DUTY_REST", "WEEKEND", "HOLIDAY", "EVIDENCE", "ABSENT"]:
                s_name = "Off / Marked"
            else:
                s_name = "Awaiting"

        out.append(AttendanceDailyOut(
            id=r.id,
            personnel_id=r.personnel_id,
            attendance_date=r.attendance_date,
            first_in=to_local(r.first_in),
            last_out=to_local(r.last_out),
            total_work_minutes=r.total_work_minutes,
            status=r.status,
            late_minutes=r.late_minutes,
            overtime_minutes=r.overtime_minutes,
            source=r.source,
            calculated_at=to_local(r.calculated_at),
            personnel_name=p.full_name if p else None,
            employee_code=p.employee_code if p else None,
            biometric_user_id=p.biometric_user_id if p else None,
            department_name=p.department.name if p and p.department else None,
            rank_name=p.rank.name if p and p.rank else None,
            course_name=p.course.name if p and p.course else None,
            category=p.category if p else None,
            designation=p.designation if p else None,
            gender=p.gender if p else None,
            is_trainee=p.is_trainee if p else None,
            cnic=p.cnic if p else None,
            father_name=p.father_name if p else None,
            shift_name=s_name,
        ))

    return PaginatedResponse(
        data=out,
        pagination=PaginationMeta(page=page, page_size=page_size, total=total),
    )



@router.get("/recent-punches", response_model=ApiResponse)
async def recent_punches(
    limit: Annotated[int, Query(ge=1, le=50)] = 10,
    db: AsyncSession = Depends(get_db),
):
    """Get the most recent punches (for live feed)."""
    result = await db.execute(
        select(AttendancePunch)
        .order_by(AttendancePunch.punch_time.desc())
        .limit(limit)
    )
    punches = result.scalars().all()
    people = (
        await db.execute(
            select(Personnel).options(
                selectinload(Personnel.rank),
                selectinload(Personnel.department),
                selectinload(Personnel.course),
            )
        )
    ).scalars().all()
    people_by_pin = {str(person.biometric_user_id): person for person in people}

    out = []
    for p in punches:
        person = _person_for_punch_pin(p.biometric_user_id, people_by_pin)
        rank_course = None
        dept_name = None
        emp_code = None
        if person:
            emp_code = person.employee_code
            dept_name = person.department.name if person.department else None
            if person.is_trainee and person.course:
                rank_course = person.course.name
            elif person.rank:
                rank_course = person.rank.name
            else:
                rank_course = person.designation

        out.append(AttendancePunchOut(
            id=p.id,
            device_id=p.device_id,
            biometric_user_id=p.biometric_user_id,
            punch_time=to_local(p.punch_time),
            punch_type=p.punch_type,
            verified=p.verified,
            source=p.source,
            created_at=to_local(p.created_at),
            personnel_name=person.full_name if person else None,
            rank_name=rank_course,
            department_name=dept_name,
            employee_code=emp_code,
            device_name=p.device.name if p.device else None,
        ))

    return ApiResponse(data=out)


@router.post("/process", response_model=ApiResponse)
@router.post("/process-daily", response_model=ApiResponse)
async def process_attendance(
    target_date: Annotated[date | None, Query(alias="date")] = None,
    personnel_id: Annotated[int | None, Query()] = None,
    db: AsyncSession = Depends(get_db),
):
    """Manually trigger attendance calculation for a specific date (defaults to today)."""
    from app.services.attendance_engine import AttendanceService
    from app.utils.timezone import today
    
    process_date = target_date or today()
    calc_service = AttendanceService()
    
    try:
        processed = await calc_service.process_daily_attendance(
            db=db, target_date=process_date, personnel_id=personnel_id
        )
        await db.commit()
        return ApiResponse(
            message=f"Processed daily attendance for {processed} personnel on {process_date}",
            data={"processed": processed, "date": process_date}
        )
    except Exception as e:
        await db.rollback()
        return ApiResponse(
            success=False,
            message=f"Failed to process attendance: {str(e)}"
        )


@router.get("/report")
async def get_attendance_report(
    is_trainee: Annotated[bool, Query()] = False,
    start_date: Annotated[date | None, Query()] = None,
    end_date: Annotated[date | None, Query()] = None,
    date_filter: Annotated[date | None, Query(alias="date")] = None,
    quick_filter: Annotated[str, Query()] = "all",
    search: Annotated[str | None, Query()] = None,
    department_id: Annotated[int | None, Query()] = None,
    rank_id: Annotated[int | None, Query()] = None,
    course_id: Annotated[int | None, Query()] = None,
    page: Annotated[int, Query(ge=1)] = 1,
    page_size: Annotated[int, Query(ge=1, le=500)] = 50,
    db: AsyncSession = Depends(get_db),
):
    """Fetch attendance report with rich filters and summaries."""
    from datetime import timedelta
    from sqlalchemy import or_
    from sqlalchemy.orm import selectinload

    target_start = start_date or date_filter or today()
    target_end = end_date or target_start
    if target_end < target_start:
        target_end = target_start

    # Personnel base query
    p_query = (
        select(Personnel)
        .options(
            selectinload(Personnel.rank),
            selectinload(Personnel.department),
            selectinload(Personnel.course),
        )
        .where(Personnel.is_trainee == is_trainee)
    )

    if department_id:
        p_query = p_query.where(Personnel.department_id == department_id)
    if rank_id:
        p_query = p_query.where(Personnel.rank_id == rank_id)
    if course_id:
        p_query = p_query.where(Personnel.course_id == course_id)
    if search:
        s = f"%{search.strip()}%"
        p_query = p_query.where(
            or_(
                Personnel.full_name.ilike(s),
                Personnel.employee_code.ilike(s),
                Personnel.biometric_user_id.ilike(s),
            )
        )

    p_query = p_query.order_by(Personnel.biometric_user_id.asc(), Personnel.id.asc())
    personnel_res = await db.execute(p_query)
    personnel_list = personnel_res.scalars().all()

    if not personnel_list:
        return {
            "data": [],
            "summary": {
                "total": 0, "present": 0, "late": 0, "absent": 0, "leave": 0, "total_hours": 0
            },
            "pagination": {
                "page": page, "page_size": page_size, "total": 0, "total_pages": 0
            }
        }

    personnel_ids = [p.id for p in personnel_list]

    # Fetch daily attendance records in range
    att_query = select(AttendanceDaily).where(
        AttendanceDaily.personnel_id.in_(personnel_ids),
        AttendanceDaily.attendance_date >= target_start,
        AttendanceDaily.attendance_date <= target_end,
    )
    att_res = await db.execute(att_query)
    all_att = att_res.scalars().all()

    att_map = {(a.personnel_id, a.attendance_date): a for a in all_att}

    items = []
    sr = 1
    is_single_day = (target_start == target_end)

    if is_single_day:
        for p in personnel_list:
            a = att_map.get((p.id, target_start))
            raw_status = (a.status if a and a.status else "").upper()
            first_in = a.first_in if a else None
            last_out = a.last_out if a else None
            total_minutes = a.total_work_minutes if a else 0.0

            if first_in:
                if raw_status == "LATE" or (a and a.late_minutes and a.late_minutes > 0):
                    status = "Late"
                else:
                    status = "Present"
            elif raw_status in ("LEAVE", "OSD", "MEDICAL", "DUTY_REST"):
                status = "Leave"
            elif raw_status == "WEEKEND":
                status = "Weekend"
            else:
                status = "Absent"

            if not _matches_quick_filter(status, quick_filter):
                continue

            rank_title = (p.rank.name if p.rank else None) or p.designation or (p.course.name if p.course else "Trainee" if p.is_trainee else "Staff")
            dept_title = (p.department.name if p.department else "General")

            first_in_local = to_local(first_in)
            last_out_local = to_local(last_out)
            check_in_str = first_in_local.strftime("%I:%M %p") if first_in_local else "—"
            check_out_str = last_out_local.strftime("%I:%M %p") if last_out_local else "—"
            hours_str = f"{total_minutes / 60:.2f}" if total_minutes and total_minutes > 0 else "—"

            items.append({
                "sr": sr,
                "id": p.id,
                "personnel_id": p.id,
                "name": p.full_name,
                "belt_no": p.employee_code or "—",
                "pin": p.biometric_user_id,
                "rank": rank_title,
                "department": dept_title,
                "status": status,
                "check_in": check_in_str,
                "check_out": check_out_str,
                "check_in_sort": first_in_local.strftime("%H:%M") if first_in_local else "99:99",
                "hours": hours_str,
                "late_minutes": a.late_minutes if a else 0,
                "date": target_start.isoformat(),
                "is_trainee": p.is_trainee,
            })
            sr += 1
    else:
        for p in personnel_list:
            if quick_filter == "summary":
                p_atts = [a for a in all_att if a.personnel_id == p.id]
                present_days = sum(1 for a in p_atts if a.first_in and a.status in ("PRESENT", "LATE"))
                late_days = sum(1 for a in p_atts if a.status == "LATE" or (a.late_minutes and a.late_minutes > 0))
                leave_days = sum(1 for a in p_atts if (a.status or "").upper() in ("LEAVE", "OSD", "MEDICAL", "DUTY_REST"))
                total_min = sum(a.total_work_minutes or 0.0 for a in p_atts)
                total_days = (target_end - target_start).days + 1
                absent_days = max(0, total_days - (present_days + leave_days))

                rank_title = (p.rank.name if p.rank else None) or p.designation or (p.course.name if p.course else "Trainee" if p.is_trainee else "Staff")
                dept_title = (p.department.name if p.department else "General")

                items.append({
                    "sr": sr,
                    "id": p.id,
                    "personnel_id": p.id,
                    "name": p.full_name,
                    "belt_no": p.employee_code or "—",
                    "pin": p.biometric_user_id,
                    "rank": rank_title,
                    "department": dept_title,
                    "status": "Summary",
                    "check_in": f"{present_days} Pres",
                    "check_out": f"{absent_days} Abs",
                    "hours": f"{total_min / 60:.2f}",
                    "late_minutes": late_days,
                    "date": f"{target_start} to {target_end}",
                    "is_trainee": p.is_trainee,
                    "summary_data": {
                        "present_days": present_days,
                        "late_days": late_days,
                        "absent_days": absent_days,
                        "leave_days": leave_days,
                    }
                })
                sr += 1
            else:
                cur_d = target_start
                while cur_d <= target_end:
                    a = att_map.get((p.id, cur_d))
                    raw_status = (a.status if a and a.status else "").upper()
                    first_in = a.first_in if a else None
                    last_out = a.last_out if a else None
                    total_minutes = a.total_work_minutes if a else 0.0

                    if first_in:
                        if raw_status == "LATE" or (a and a.late_minutes and a.late_minutes > 0):
                            status = "Late"
                        else:
                            status = "Present"
                    elif raw_status in ("LEAVE", "OSD", "MEDICAL", "DUTY_REST"):
                        status = "Leave"
                    elif raw_status == "WEEKEND":
                        status = "Weekend"
                    else:
                        status = "Absent"

                    if not _matches_quick_filter(status, quick_filter):
                        cur_d += timedelta(days=1)
                        continue

                    rank_title = (p.rank.name if p.rank else None) or p.designation or (p.course.name if p.course else "Trainee" if p.is_trainee else "Staff")
                    dept_title = (p.department.name if p.department else "General")

                    first_in_local = to_local(first_in)
                    last_out_local = to_local(last_out)
                    check_in_str = first_in_local.strftime("%I:%M %p") if first_in_local else "—"
                    check_out_str = last_out_local.strftime("%I:%M %p") if last_out_local else "—"
                    hours_str = f"{total_minutes / 60:.2f}" if total_minutes and total_minutes > 0 else "—"

                    items.append({
                        "sr": sr,
                        "id": p.id,
                        "personnel_id": p.id,
                        "name": p.full_name,
                        "belt_no": p.employee_code or "—",
                        "pin": p.biometric_user_id,
                        "rank": rank_title,
                        "department": dept_title,
                        "status": status,
                        "check_in": check_in_str,
                        "check_out": check_out_str,
                        "check_in_sort": first_in_local.strftime("%H:%M") if first_in_local else "99:99",
                        "hours": hours_str,
                        "late_minutes": a.late_minutes if a else 0,
                        "date": cur_d.isoformat(),
                        "is_trainee": p.is_trainee,
                    })
                    sr += 1
                    cur_d += timedelta(days=1)

    if quick_filter == "by_department":
        items.sort(key=lambda x: (x["department"], x["name"]))
    elif quick_filter == "present":
        items.sort(key=lambda x: (x.get("check_in_sort") or "99:99", x["name"]))
    elif quick_filter == "all":
        items.sort(key=lambda x: (_STATUS_SORT.get(x["status"], 9), x["name"]))
    for idx, item in enumerate(items, 1):
        item["sr"] = idx

    total_matching = len(items)
    present_cnt = sum(1 for it in items if it["status"] == "Present")
    late_cnt = sum(1 for it in items if it["status"] == "Late")
    leave_cnt = sum(1 for it in items if it["status"] == "Leave")
    absent_cnt = sum(1 for it in items if it["status"] in ("Absent", "Weekend"))
    total_hrs = sum(float(it["hours"]) for it in items if it["hours"] != "—")

    start_idx = (page - 1) * page_size
    end_idx = start_idx + page_size
    page_data = items[start_idx:end_idx]

    return {
        "data": page_data,
        "summary": {
            "total": total_matching,
            "present": present_cnt,
            "late": late_cnt,
            "absent": absent_cnt,
            "leave": leave_cnt,
            "total_hours": round(total_hrs, 2),
        },
        "pagination": {
            "page": page,
            "page_size": page_size,
            "total": total_matching,
            "total_pages": (total_matching + page_size - 1) // page_size if page_size > 0 else 1,
        }
    }


@router.get("/report/export")
async def export_attendance_report(
    is_trainee: Annotated[bool, Query()] = False,
    start_date: Annotated[date | None, Query()] = None,
    end_date: Annotated[date | None, Query()] = None,
    date_filter: Annotated[date | None, Query(alias="date")] = None,
    quick_filter: Annotated[str, Query()] = "all",
    search: Annotated[str | None, Query()] = None,
    department_id: Annotated[int | None, Query()] = None,
    rank_id: Annotated[int | None, Query()] = None,
    course_id: Annotated[int | None, Query()] = None,
    format: Annotated[str, Query()] = "xlsx",
    db: AsyncSession = Depends(get_db),
):
    """Export attendance report as Excel (.xlsx) or CSV."""
    import io
    import csv
    from fastapi.responses import Response

    report = await get_attendance_report(
        is_trainee=is_trainee,
        start_date=start_date,
        end_date=end_date,
        date_filter=date_filter,
        quick_filter=quick_filter,
        search=search,
        department_id=department_id,
        rank_id=rank_id,
        course_id=course_id,
        page=1,
        page_size=10000,
        db=db,
    )
    rows = report.get("data", [])
    target_date_str = str(start_date or date_filter or today())
    entity_label = "Trainees" if is_trainee else "Staff"

    if format.lower() == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["SR", "NAME", "BELT NO", "PIN", "RANK / COURSE", "DEPARTMENT", "STATUS", "CHECK-IN", "CHECK-OUT", "HOURS", "DATE"])
        for r in rows:
            writer.writerow([
                r["sr"],
                r["name"],
                r["belt_no"],
                r["pin"],
                r["rank"],
                r["department"],
                r["status"],
                r["check_in"],
                r["check_out"],
                r["hours"],
                r["date"],
            ])
        csv_bytes = output.getvalue().encode("utf-8-sig")
        return Response(
            content=csv_bytes,
            media_type="text/csv",
            headers={
                "Content-Disposition": f'attachment; filename="Attendance_{entity_label}_{target_date_str}.csv"'
            },
        )

    # Excel format
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{entity_label} Attendance"
    ws.views.sheetView[0].showGridLines = True

    # Title header
    setting_res = await db.execute(select(SystemSetting.value).where(SystemSetting.key == "org_display_name"))
    org_display_name = setting_res.scalar_one_or_none() or "Police Training School Sargodha"

    ws.merge_cells("A1:K1")
    ws["A1"] = f"{org_display_name} — {entity_label} Attendance Report ({target_date_str})"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ["SR", "NAME", "BELT NO", "PIN", "RANK / COURSE", "DEPARTMENT", "STATUS", "CHECK-IN", "CHECK-OUT", "HOURS", "DATE"]
    ws.append([])
    ws.append(headers)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[3].height = 24

    for r in rows:
        row_data = [
            r["sr"],
            r["name"],
            r["belt_no"],
            r["pin"],
            r["rank"],
            r["department"],
            r["status"],
            r["check_in"],
            r["check_out"],
            r["hours"],
            r["date"],
        ]
        ws.append(row_data)
        cur_row = ws.max_row
        ws.row_dimensions[cur_row].height = 20
        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=cur_row, column=col_num)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if col_num in (1, 3, 4, 7, 8, 9, 10, 11):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Status color highlights
            if col_num == 7:
                st = str(cell.value).upper()
                if "PRESENT" in st:
                    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="166534")
                elif "LATE" in st:
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="92400E")
                elif "LEAVE" in st:
                    cell.fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="075985")
                elif "ABSENT" in st:
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B")

    # Column width auto-fit
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="Attendance_{entity_label}_{target_date_str}.xlsx"'
        },
    )


# ════════════════════════════════════════════════════════════════
# Exceptions (Leave, OSD, Medical, Duty Rest, Repatriation)
# ════════════════════════════════════════════════════════════════

class ExceptionCreate(BaseModel):
    personnel_id: int
    date: dt_date | None = None
    start_date: dt_date | None = None
    end_date: dt_date | None = None
    exception_type: str  # LEAVE, OSD, MEDICAL, DUTY_REST, REPATRIATION, EVIDENCE, PRESENT, ABSENT
    reason: str | None = None
    approved_by: str | None = None
    remarks: str | None = None


@router.post("/exceptions", response_model=ApiResponse)
async def create_or_update_exception(
    payload: ExceptionCreate,
    db: AsyncSession = Depends(get_db),
):
    """Create or update an attendance exception (Leave, OSD, Medical, Duty Rest, Repatriation) across single date or range."""
    p_res = await db.execute(select(Personnel).where(Personnel.id == payload.personnel_id))
    person = p_res.scalar_one_or_none()
    if not person:
        raise HTTPException(status_code=404, detail="Personnel record not found")

    s_date = payload.start_date or payload.date or today()
    e_date = payload.end_date or s_date
    if s_date > e_date:
        s_date, e_date = e_date, s_date

    current = s_date
    calc_service = AttendanceService()
    while current <= e_date:
        existing_res = await db.execute(
            select(AttendanceException).where(
                AttendanceException.personnel_id == payload.personnel_id,
                AttendanceException.date == current,
            )
        )
        existing = existing_res.scalar_one_or_none()

        if existing:
            existing.exception_type = payload.exception_type.upper()
            existing.reason = payload.reason
            existing.approved_by = payload.approved_by
            existing.remarks = payload.remarks
            existing.updated_at = now()
        else:
            new_exc = AttendanceException(
                personnel_id=payload.personnel_id,
                date=current,
                exception_type=payload.exception_type.upper(),
                reason=payload.reason,
                approved_by=payload.approved_by,
                remarks=payload.remarks,
            )
            db.add(new_exc)

        await db.flush()
        await calc_service.process_daily_attendance(db, current, personnel_id=payload.personnel_id)
        current += timedelta(days=1)

    await db.commit()

    return ApiResponse(
        message=f"Attendance exception '{payload.exception_type.upper()}' recorded for {person.full_name} from {s_date} to {e_date}",
        data={"personnel_id": payload.personnel_id, "start_date": str(s_date), "end_date": str(e_date), "status": payload.exception_type.upper()},
    )


@router.get("/exceptions", response_model=ApiResponse)
async def list_exceptions(
    date_filter: Annotated[date | None, Query(alias="date")] = None,
    personnel_id: Annotated[int | None, Query()] = None,
    exception_type: Annotated[str | None, Query()] = None,
    db: AsyncSession = Depends(get_db),
):
    """List attendance exceptions with filters."""
    query = select(AttendanceException).options(
        selectinload(AttendanceException.personnel).selectinload(Personnel.rank),
        selectinload(AttendanceException.personnel).selectinload(Personnel.department),
        selectinload(AttendanceException.personnel).selectinload(Personnel.course),
    ).order_by(AttendanceException.date.desc())

    if date_filter:
        query = query.where(AttendanceException.date == date_filter)
    if personnel_id:
        query = query.where(AttendanceException.personnel_id == personnel_id)
    if exception_type:
        query = query.where(AttendanceException.exception_type == exception_type.upper())

    result = await db.execute(query)
    exceptions = result.scalars().all()

    out = []
    for exc in exceptions:
        p = exc.personnel
        rank_or_course = None
        dept_name = None
        if p:
            dept_name = p.department.name if p.department else None
            rank_or_course = p.course.name if (p.is_trainee and p.course) else (p.rank.name if p.rank else p.designation)

        out.append({
            "id": exc.id,
            "personnel_id": exc.personnel_id,
            "personnel_name": p.full_name if p else "N/A",
            "biometric_user_id": p.biometric_user_id if p else "N/A",
            "employee_code": p.employee_code if p else None,
            "rank_or_course": rank_or_course,
            "department_name": dept_name,
            "is_trainee": p.is_trainee if p else False,
            "date": str(exc.date),
            "exception_type": exc.exception_type,
            "reason": exc.reason,
            "approved_by": exc.approved_by,
            "remarks": exc.remarks,
            "created_at": exc.created_at.isoformat() if exc.created_at else None,
        })

    return ApiResponse(data=out)


@router.delete("/exceptions/{exception_id}", response_model=ApiResponse)
async def delete_exception(
    exception_id: int,
    db: AsyncSession = Depends(get_db),
):
    """Delete an exception and recalculate the day's attendance."""
    result = await db.execute(select(AttendanceException).where(AttendanceException.id == exception_id))
    exc = result.scalar_one_or_none()
    if not exc:
        raise HTTPException(status_code=404, detail="Exception not found")

    target_date = exc.date
    target_personnel_id = exc.personnel_id

    await db.delete(exc)
    await db.flush()

    calc_service = AttendanceService()
    await calc_service.process_daily_attendance(db, target_date, personnel_id=target_personnel_id)
    await db.commit()

    return ApiResponse(message="Exception deleted and attendance recalculated")


# ════════════════════════════════════════════════════════════════
# Manual Punch
# ════════════════════════════════════════════════════════════════

class ManualPunchCreate(BaseModel):
    personnel_id: int | None = None
    biometric_user_id: str | None = None
    punch_time: datetime
    punch_type: str | None = None
    punch_state: int | None = None
    device_id: int | None = None
    remarks: str | None = None


@router.post("/manual-punch", response_model=ApiResponse)
async def record_manual_punch(
    payload: ManualPunchCreate,
    db: AsyncSession = Depends(get_db),
):
    """Record a manual biometric punch and recalculate attendance for that day."""
    bio_id = payload.biometric_user_id
    person = None
    if payload.personnel_id:
        p_res = await db.execute(select(Personnel).where(Personnel.id == payload.personnel_id))
        person = p_res.scalar_one_or_none()
        if person:
            bio_id = person.biometric_user_id
    elif bio_id:
        p_res = await db.execute(select(Personnel).where(Personnel.biometric_user_id == bio_id))
        person = p_res.scalar_one_or_none()

    if not bio_id:
        raise HTTPException(status_code=400, detail="Must provide personnel_id or biometric_user_id")

    dev_id = payload.device_id
    if not dev_id:
        d_res = await db.execute(select(Device).where(Device.enabled == True).limit(1))
        dev = d_res.scalar_one_or_none()
        dev_id = dev.id if dev else 1

    # Resolve punch type: support punch_state (1 = IN, 2 = OUT) or punch_type ("IN", "OUT")
    resolved_type = "IN"
    if payload.punch_state == 2 or (payload.punch_type and payload.punch_type.upper() == "OUT"):
        resolved_type = "OUT"
    elif payload.punch_state == 1 or (payload.punch_type and payload.punch_type.upper() == "IN"):
        resolved_type = "IN"
    elif payload.punch_type:
        resolved_type = payload.punch_type.upper()

    target_time = payload.punch_time

    # Check for duplicate punch at the exact same device, user, and timestamp
    existing_res = await db.execute(
        select(AttendancePunch).where(
            AttendancePunch.device_id == dev_id,
            AttendancePunch.biometric_user_id == bio_id,
            AttendancePunch.punch_time == target_time,
        )
    )
    existing_punch = existing_res.scalar_one_or_none()

    if existing_punch:
        if existing_punch.punch_type == resolved_type:
            # Exact duplicate already recorded
            punch = existing_punch
        else:
            # Different punch state at same minute (e.g. IN followed immediately by OUT)
            # Offset by 1 second to avoid UniqueConstraint collision
            target_time = target_time + timedelta(seconds=1)
            punch = AttendancePunch(
                device_id=dev_id,
                biometric_user_id=bio_id,
                punch_time=target_time,
                punch_type=resolved_type,
                source="MANUAL",
                verified=1,
            )
            db.add(punch)
    else:
        punch = AttendancePunch(
            device_id=dev_id,
            biometric_user_id=bio_id,
            punch_time=target_time,
            punch_type=resolved_type,
            source="MANUAL",
            verified=1,
        )
        db.add(punch)

    try:
        await db.flush()
    except Exception as exc:
        logger.warning("Integrity conflict while saving punch, applying fallback offset: %s", exc)
        await db.rollback()
        target_time = payload.punch_time + timedelta(seconds=1)
        punch = AttendancePunch(
            device_id=dev_id,
            biometric_user_id=bio_id,
            punch_time=target_time,
            punch_type=resolved_type,
            source="MANUAL",
            verified=1,
        )
        db.add(punch)
        await db.flush()

    punch_date = target_time.date()
    calc_service = AttendanceService()
    p_id = person.id if person else None
    await calc_service.process_daily_attendance(db, punch_date, personnel_id=p_id)
    await db.commit()

    return ApiResponse(
        message=f"Manual {resolved_type} punch recorded for user {bio_id} at {target_time.strftime('%Y-%m-%d %H:%M:%S')}",
        data={"biometric_user_id": bio_id, "punch_time": target_time.isoformat(), "punch_type": resolved_type},
    )


# ════════════════════════════════════════════════════════════════
# Unlinked Punches
# ════════════════════════════════════════════════════════════════

@router.get("/unlinked", response_model=ApiResponse)
async def list_unlinked_punches(db: AsyncSession = Depends(get_db)):
    """List biometric punches that do not match any registered personnel."""
    registered_ids = select(Personnel.biometric_user_id)

    query = (
        select(
            AttendancePunch.biometric_user_id,
            func.count(AttendancePunch.id).label("punch_count"),
            func.max(AttendancePunch.punch_time).label("last_punch"),
            func.min(AttendancePunch.device_id).label("device_id"),
        )
        .where(AttendancePunch.biometric_user_id.not_in(registered_ids))
        .group_by(AttendancePunch.biometric_user_id)
        .order_by(func.max(AttendancePunch.punch_time).desc())
    )

    result = await db.execute(query)
    rows = result.all()

    out = []
    for r in rows:
        dev_res = await db.execute(select(Device).where(Device.id == r.device_id))
        dev = dev_res.scalar_one_or_none()
        lp_local = to_local(r.last_punch)
        out.append({
            "pin": str(r.biometric_user_id),
            "punch_count": r.punch_count,
            "last_punch": lp_local.strftime("%d %b %I:%M %p") if lp_local else "N/A",
            "last_punch_iso": lp_local.isoformat() if lp_local else None,
            "device": dev.name if dev else f"Device #{r.device_id}",
        })

    return ApiResponse(data=out)

